"""Build docs/PotholePal-day2-backend.docx.

Re-run after editing the content below:  python docs/build_day2_docx.py
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

BLUE = RGBColor(0x1D, 0x5B, 0xEB)
GREY = RGBColor(0x64, 0x74, 0x8B)

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(10)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def note(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = GREY
    run.font.size = Pt(10)


def code(lines):
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)


# ---------------------------------------------------------------- title
title = doc.add_heading("PotholePal: day two, the backend", level=0)
title.runs[0].font.color.rgb = BLUE
sub = doc.add_paragraph()
run = sub.add_run("Tuesday 20 October 2026  ·  The City Fix Team (Lesiamo, Crystal, Kgosi, Raymond)")
run.font.color.rgb = GREY
run.font.size = Pt(10)

doc.add_paragraph(
    "Yesterday the app looked right. Today it became a system: real logins, a real "
    "database, reports that survive a reload, and a staff portal that changes what the "
    "resident sees. Everything in this document was run and tested on this machine "
    "today, so you can repeat it and you can explain it."
)

# ---------------------------------------------------------------- 1
doc.add_heading("Where the system stands now", level=1)
doc.add_paragraph(
    "A resident signs in, submits a pothole with a photo, gets a reference number from "
    "the database, and watches the status move along the timeline. A staff member signs "
    "in through a separate portal, sees the same reports sorted by priority, and advances "
    "the status. The resident's screen picks the change up within ten seconds without a "
    "reload. Two browser profiles were used to prove it today; on Thursday use Chrome for "
    "one and Edge for the other."
)
table(
    ["Piece", "What it is", "Where it lives"],
    [
        ["Front end", "React 19.3, Vite 7.3, Tailwind 4, Leaflet 1.9, react-router 7.18", "src/, built into dist/"],
        ["API server", "Express 5.3, 14 endpoints, bearer-token sessions", "server/index.js, port 3001"],
        ["Database", "SQL Server 2022, database PotholePal, 4 tables", "your local instance MSSQLSERVER"],
        ["Photos", "Uploaded files on disk, path stored in the reports table", "server/uploads/ (git-ignored)"],
        ["Passwords", "bcrypt hashes, 10 rounds", "users.password_hash"],
        ["Secrets", "Database login password", "server/.env (git-ignored)"],
    ],
    widths=[1.2, 3.4, 2.0],
)
note(
    "The day-one handover suggested SQLite as a fallback. SQL Server was chosen instead, "
    "as recorded below. Where the two documents disagree on stack, this one is current."
)

# ---------------------------------------------------------------- 2
doc.add_heading("The decisions you asked about", level=1)

doc.add_heading("Database: SQL Server", level=2)
doc.add_paragraph(
    "The rubric names no database. Criterion A3 asks for a well-justified stack, which is "
    "about the defence rather than the choice. SQL Server was already installed on this "
    "machine (2022, default instance, SSMS 22, your class databases on the same instance), "
    "your group uses it in class, and it needs no install before the demo. MySQL is present "
    "through C:\\xampp and would have worked, but switching would have gone against the "
    "group decision for nothing. SQLite installs fastest and then leaves you explaining in "
    "the viva why the demo runs on a database the group did not choose."
)
doc.add_paragraph(
    "If the assessor asks, the short answer is: we chose SQL Server because it is what we "
    "use in class, it was already available on our machines, and the rubric asks us to "
    "justify a choice rather than to use a particular product."
)

doc.add_heading("Hosting: run it on the laptop", level=2)
doc.add_paragraph(
    "InfinityFree is PHP and MySQL hosting. It cannot run a Node.js backend, so hosting "
    "there and keeping the Express server cannot both happen. A search for free hosting "
    "covering a Node API together with a SQL database returned nothing that covers both. "
    "For Thursday 29 October the backend runs on the laptop: both browsers open "
    "localhost, the classroom needs no internet connection, and nothing can fall asleep "
    "mid-demo. If you still want a public URL for the dissertation after Tuesday, that is "
    "the moment to look again."
)

doc.add_heading("Accounts: seeded, simple, hashed", level=2)
doc.add_paragraph(
    "Five accounts exist in the database for testing. Login is email and password, which "
    "is what you asked for. The passwords are hashed with bcrypt rather than stored in "
    "plain text, because that is one assessor question you can then answer with evidence "
    "instead of a promise. Registration, password reset and rate limiting are not built "
    "yet; security hardening stays on the plan for after the demo."
)

doc.add_heading("Staff portal: a separate login", level=2)
doc.add_paragraph(
    "The portal sits at /staff with its own login screen and its own stored session, and "
    "it talks to the same API and the same database as the resident app. The Profile "
    "button that used to say demo now says staff and points there. Sessions are stored "
    "under separate keys, so one browser can even hold both at once, while two different "
    "browsers stay completely independent. That is the test you described, and it passed "
    "today."
)

doc.add_heading("Dates in play", level=2)
bullets([
    "Application ready for testing: tonight, Wednesday 21 October at the latest, as you planned.",
    "Triage comparison figures: Thursday 22 October. The priority scores in the database feed them directly.",
    "Rehearsal and presentation preparation: Monday 26 October.",
    "Dissertation and demo submission: Thursday 29 October, Turnitin similarity below 15 percent.",
])

# ---------------------------------------------------------------- 3
doc.add_heading("What is in the database", level=1)
table(
    ["Table", "What it holds", "Worth knowing"],
    [
        ["users", "5 accounts: name, email, password_hash, role, points, level", "role is resident or staff; email is unique"],
        ["reports", "One row per work item: location, size, depth, traffic, damage, status, priority, score, photo path, evidence path", "Identity starts at 1018, so live reports continue from #1025"],
        ["report_status_history", "Every status change, who changed it, when", "Append only; the details timeline is built from it"],
        ["notifications", "Messages for the resident: received, verified, scheduled, fixed, points", "Created by the server when something happens"],
    ],
    widths=[1.3, 3.3, 2.0],
)

doc.add_heading("The reference number trick", level=2)
doc.add_paragraph(
    "The reports identity starts at 1018. The seven designed seed reports fill 1018 to "
    "1024, so the first report anyone submits in the demo comes out as #1025, which is the "
    "number the day-one document used. The 90 older synthetic rows that give the portal "
    "some volume carry explicit ids 1 to 90. Because those sit below 1018 they do not move "
    "the identity counter, and the next live report still gets 1025."
)
doc.add_paragraph(
    "The reference column is a persisted computed column: the hash sign followed by the "
    "id, with a unique index on it. No application code ever generates a reference number, "
    "so two reports can never share one."
)

doc.add_heading("Two changes made to SQL Server itself", level=2)
doc.add_paragraph(
    "The instance was accepting Windows logins only, which is why the application login "
    "was rejected at first, and its TCP protocol was switched off, which is why the SQL "
    "tools worked while the Node driver could not connect. Both are normal settings on a "
    "class laptop, and both had to change for the backend to connect:"
)
bullets([
    "Mixed-mode authentication enabled (registry value LoginMode set to 2, service restarted). "
    "In SSMS the same change is Server Properties, Security page, SQL Server and Windows Authentication mode.",
    "TCP/IP enabled on the two loopback addresses only, static port 1433 (service restarted again). "
    "In SQL Server Configuration Manager the same change is Network Configuration, Protocols, TCP/IP, "
    "with port 1433 on the IP addresses used. Loopback only, so nothing opened to the network; the firewall was not touched.",
])
doc.add_paragraph(
    "The script that made both changes is saved at server/db/enable-tcp.ps1. On a new "
    "machine you need the same two settings, through SSMS or through that script, before "
    "the backend will connect."
)
doc.add_paragraph(
    "The application does not connect as an administrator. The login potholepal owns only "
    "the PotholePal database, and its password sits in server/.env, which is git-ignored."
)

# ---------------------------------------------------------------- 4
doc.add_heading("The server, file by file", level=1)
table(
    ["Path", "What it does"],
    [
        ["server/index.js", "The Express app: all 14 endpoints, session tokens, upload handling, static serving of uploads and dist/."],
        ["server/db/schema.sql", "The four tables and their indexes, written so it can run more than once without errors."],
        ["server/db/setup.js", "Runs the schema. npm run db:setup."],
        ["server/db/seed.js", "Creates the five accounts, the seven designed reports with their history, 90 synthetic older reports, and six notifications. npm run db:seed."],
        ["server/db/reset.js", "Deletes everything, reseeds the identity, clears uploaded photos. Runs seed after itself via npm run db:reset."],
        ["server/lib/domain.js", "The rules: priority scoring, severity labels, nearest area, distance, time formatting. Shared by the API and the seeder, and it reads the area coordinates from src/lib/data.js so both sides agree."],
        ["server/db/db.js", "The connection pool, configured from server/.env."],
        ["server/smoke-test.sh", "26 automated checks against a running server. bash server/smoke-test.sh."],
        ["server/.env", "DB_SERVER, DB_NAME, DB_USER, DB_PASSWORD, DB_PORT. Git-ignored."],
    ],
    widths=[1.7, 4.9],
)

doc.add_heading("The endpoints that actually exist", level=2)
table(
    ["Method", "Path", "Who", "What it does"],
    [
        ["POST", "/api/v1/auth/login", "public", "Email and password, returns a token and the user."],
        ["POST", "/api/v1/auth/logout", "any", "Drops the token."],
        ["GET", "/api/v1/auth/me", "any", "The signed-in user with live points, for Profile."],
        ["GET", "/api/v1/reports", "any", "Every report with reporter name, timeline, isMine, priority and score."],
        ["GET", "/api/v1/reports/:id", "any", "One report, for the details screen when a session opens it directly."],
        ["GET", "/api/v1/reports/:id/history", "any", "Status changes in order."],
        ["POST", "/api/v1/reports", "resident", "Multipart: photo file plus location and issue fields. Runs the duplicate check, scores the priority, writes report, history and notification, returns the report."],
        ["POST", "/api/v1/reports/:id/evidence", "resident", "Attaches a supporting photo to an existing report instead of creating a second one."],
        ["PATCH", "/api/v1/reports/:id", "staff", "Changes status, stamps scheduled_at or fixed_at, appends history, notifies the reporter, awards 50 points on fixed."],
        ["GET", "/api/v1/staff/queue", "staff", "Reports ordered by priority score, for the portal."],
        ["GET", "/api/v1/staff/stats", "staff", "Totals, new count, fixed count, resolution rate, heatmap points, recent reports."],
        ["GET", "/api/v1/stats", "resident", "Home screen: nearby, total, fixed."],
        ["GET", "/api/v1/notifications", "resident", "Grouped Today and Earlier, matching the screen."],
        ["POST", "/api/v1/notifications/:id/read", "resident", "Marks one notification read."],
    ],
    widths=[0.7, 2.1, 0.85, 3.05],
)
note(
    "Registration from the day-one plan was left out on purpose: accounts are seeded, "
    "which is what you asked for. Everything else from that plan's endpoint table is live."
)

doc.add_heading("The rules, as implemented", level=2)
doc.add_paragraph(
    "Priority scoring runs on every submission and the score is stored, because the "
    "triage comparison figures due Thursday 22 October need the raw numbers:"
)
table(
    ["Field", "Points"],
    [
        ["size", "Small 1, Medium 2, Large 3"],
        ["depth", "Shallow 1, Moderate 2, Deep 3"],
        ["traffic", "Low 1, Medium 2, High 3"],
        ["damage", "No 0, Yes 2"],
    ],
    widths=[1.3, 4.9],
)
doc.add_paragraph(
    "P1 at 9 or more, or damage reported on a high traffic road. P2 from 6 to 8. P3 at 5 "
    "or below. The portal queue is sorted by this score, so the worst reports arrive at "
    "the top. In today's test a Large, Deep, High traffic report with damage scored 11 "
    "and opened at position one."
)
doc.add_paragraph(
    "Duplicate matching: on submission the server looks for a report that is not fixed, "
    "within 50 metres, created in the last 7 days. It finds candidates with a bounding "
    "box query and then measures exact distance. If one matches, no second work item is "
    "created: the response is 409 with the existing reference, that report's duplicate "
    "count goes up, and the screen offers to attach your photo as supporting evidence."
)

# ---------------------------------------------------------------- 5
doc.add_heading("What changed in the front end", level=1)
table(
    ["Screen", "What changed today"],
    [
        ["src/lib/api.js (new)", "Fetch wrapper: sends the bearer token, turns error envelopes into exceptions, and on a 401 clears the session and returns to the right login screen."],
        ["Login (new)", "Resident sign-in with the demo account printed on the screen."],
        ["StaffLogin (new)", "Portal sign-in on the navy background, with its own demo account."],
        ["App.jsx", "Route guards for resident pages and the details page, reports loaded from the API, a refresh every 10 seconds, and submitDraft turned into a real multipart POST."],
        ["ReportPotholePage4", "Submit is now async with a busy state, and a failed duplicate shows the 409 message with an add-my-photo-as-evidence button."],
        ["Home", "Stats come from /stats, markers come from the live list, and the fake cluster pins are gone because 97 real markers exist now."],
        ["Map, MyReports", "Both read the live list from context; My Reports filters on isMine, which the server computes from the logged-in user."],
        ["ReportDetails", "Reporter name comes from the database, timeline dates come from the history table, and a staff session sees a Staff actions bar with the priority badge and an advance button."],
        ["Notifications", "Groups fetched from the API; the six seeded messages and any new ones arrive from the database."],
        ["Dashboard (portal)", "Real stat tiles, real heatmap points, recent reports from the API, queue sorted by score, and the exit button now signs the staff session out."],
        ["Profile", "Name, points and level come from the signed-in account, the report count is the real count, and there is a sign out button."],
    ],
    widths=[1.5, 5.1],
)

# ---------------------------------------------------------------- 6
doc.add_heading("How a request travels", level=1)
doc.add_paragraph("Submitting a report, step by step, which is the flow to walk through in the viva:")
numbered([
    "The review screen builds a FormData body: the photo as a file, plus lat, lng, address, issue type, size, depth, traffic, damage and description.",
    "The browser posts it to /api/v1/reports with the resident's bearer token. In development Vite proxies /api to port 3001, so the page still sees one origin.",
    "Express checks the session, then multer saves the photo into server/uploads under a generated name.",
    "The duplicate check runs: bounding box in SQL, exact distance in JavaScript, 50 metres and 7 days and not fixed.",
    "The scoring rule produces a score and a priority. The nearest area is worked out from the seven area coordinates, which is why the review screen could show Three Rivers without anyone typing it.",
    "One insert into reports, one history row for the new status, one notification for the reporter. The identity gives 1025 and the computed column makes it #1025.",
    "The report comes back as JSON in the shape the screens already expected, the draft clears, and the details page opens on the new reference.",
])
doc.add_paragraph(
    "The two-browser sync is simpler than it looks. Nothing pushes data anywhere. Both "
    "sessions poll GET /api/v1/reports every 10 seconds, and both read the same rows in "
    "the same table. When the staff browser writes a status change, the resident's next "
    "poll returns it. In today's test the staff tab moved a report to Verified at 11:38 "
    "and the resident tab, untouched, showed the same timestamp seconds later."
)

# ---------------------------------------------------------------- 7
doc.add_heading("Running it", level=1)
doc.add_heading("Development, two terminals", level=2)
code([
    "cd server && npm start        # API on http://localhost:3001",
    "npm run dev                   # front end on http://localhost:5173, proxies /api",
])
doc.add_heading("Demo mode, one terminal, one port (verified today)", level=2)
code([
    "npm run build                 # builds dist/",
    "cd server && npm start        # everything on http://localhost:3001",
])
doc.add_paragraph(
    "In demo mode the server hands out the built front end itself, and deep links such as "
    "/login still resolve. Both were checked today: the root and a route both answered 200 "
    "with HTML. One process to start on Thursday, one fewer thing to go wrong."
)
doc.add_heading("Data commands", level=2)
code([
    "npm run db:setup     # create the tables, safe to repeat",
    "npm run db:seed      # load accounts, reports, notifications, safe to repeat",
    "npm run db:reset     # wipe everything and seed again",
])
doc.add_paragraph(
    "Use db:reset between practice runs. It was tested today: after a full reset and seed "
    "the database reads 97 reports, designed ids 1018 to 1024, identity at 1024, so the "
    "next submission in your practice demo is #1025 again."
)
doc.add_heading("Accounts", level=2)
table(
    ["Role", "Email", "Password", "Who it is"],
    [
        ["Resident", "alex@potholepal.test", "resident123", "Alex R., the demo reporter"],
        ["Resident", "john@potholepal.test", "resident123", "Owns some seeded reports"],
        ["Resident", "thabo@potholepal.test", "resident123", "Owns some seeded reports"],
        ["Resident", "sarah@potholepal.test", "resident123", "Owns some seeded reports"],
        ["Staff", "staff@potholepal.test", "staff123", "P. Mokoena, the portal"],
    ],
    widths=[0.9, 2.1, 1.3, 2.3],
)
doc.add_heading("The smoke test", level=2)
doc.add_paragraph(
    "With the server running: bash server/smoke-test.sh. It runs 26 checks and prints "
    "PASS or FAIL for each, exiting 0 only when everything passed. It covers signed-out "
    "401s, both logins, a wrong password, the report list and ownership flags, the "
    "stats and notification counts, a real submission landing on #1025 with score 11 and "
    "P1, a duplicate returning 409 with that reference, a resident being refused a status "
    "change with 403, a staff status change with its timeline entry, the notifications that "
    "change creates, queue ordering by score, staff statistics, the evidence attachment, "
    "and the count staying at one work item. Today's run finished with all 26 passing. "
    "Keep that output: criterion A4 wants empirical testing, and this is it."
)
note(
    "The smoke test submits a real report, so it consumes #1025. Run npm run db:reset "
    "afterwards and the demo is clean again."
)

# ---------------------------------------------------------------- run it yourself
doc.add_heading("Run it yourself, in the browser", level=1)
doc.add_paragraph(
    "The whole procedure from a cold machine, written out so you can do it without "
    "asking anyone."
)
numbered([
    "Open a terminal in the project folder, the one that contains package.json. "
    "Check Node first: node -v should print v24.22.0.",
    "Start the API: cd server, then npm start. Wait for the line PotholePal API on "
    "http://localhost:3001 (db: PotholePal). Leave that window open; Ctrl+C stops the app.",
    "Make sure there is a current build: npm run build in the project folder, any time "
    "the code changed.",
    "Open http://localhost:3001 in a browser. In this mode one process serves the app "
    "and the API together, and that is the setup for Thursday.",
    "For editing with hot reload instead, use two windows: npm start inside server/, and "
    "npm run dev in the project folder, then open http://localhost:5173. Vite forwards "
    "/api and /uploads to port 3001.",
    "The first screen is the login, with the demo credentials printed on it. Sign in as "
    "alex@potholepal.test / resident123.",
    "Walk the app: Report a Pothole (the sample photo button works on a laptop), confirm "
    "the location, pick the chips, submit, and the reference number comes back from the "
    "database. Then the map, notifications, profile, and the municipality portal button.",
    "For the two-sided demo, open two browser windows: Chrome signed in as the resident, "
    "Edge signed in as staff@potholepal.test / staff123.",
    "Between practice runs: cd server, then npm run db:reset. The data starts clean and "
    "the next submitted report is #1025 again.",
])
doc.add_paragraph("When something does not start:")
bullets([
    "The login says the request failed: the API window is not running. Start it (step 2).",
    "Port 3001 is busy: netstat -ano | findstr 3001 shows the process id, then "
    "taskkill /PID that-id /F.",
    "A details screen says could not be found: that report id no longer exists after a "
    "reset. Go back to Home.",
    "npm is not found: that terminal has no Node on its PATH. Node lives at "
    "C:\\Program Files\\nodejs; close the terminal and open a new one.",
])

# ---------------------------------------------------------------- phone
doc.add_heading("Testing it on your phone", level=1)
doc.add_paragraph(
    "There is no executable file to copy onto a phone, and the reason is worth "
    "understanding: the app is a front end plus a server plus SQL Server, and those run "
    "on the laptop. The phone joins as a second screen for the same system over your "
    "Wi-Fi. The server already listens on every network interface, so what is left is "
    "the same Wi-Fi and one firewall rule."
)
numbered([
    "Put the laptop and the phone on the same Wi-Fi network.",
    "Start the demo on the laptop: npm run build, then cd server && npm start, and leave "
    "it running.",
    "Find the laptop's address: ipconfig, then the IPv4 address under the Wi-Fi adapter. "
    "Today it is 192.168.1.221. It changes when you move to another network.",
    "One-time firewall step, in a terminal opened as administrator. Until this has run, "
    "Windows refuses the phone, not the app:",
])
code([
    'netsh advfirewall firewall add rule name="PotholePal demo 3001" dir=in action=allow protocol=TCP localport=3001',
])
doc.add_paragraph("Then, on the phone:")
numbered([
    "Open http://192.168.1.221:3001 in the phone browser: the laptop's address, the same port.",
    "Sign in with the demo account printed on the login screen.",
    "Add an icon: iPhone, Safari, Share, Add to Home Screen; Android, Chrome menu, Add to "
    "Home screen. Either opens the app full screen from your home screen. The manifest "
    "and the 192 and 512 pixel icons for that are already in place.",
])
doc.add_paragraph(
    "A real .apk is a different project. It needs Android Studio, a current JDK and "
    "Gradle, none of which are on this machine, so roughly a gigabyte of tools and a day "
    "of work. Even built, it would still call this laptop for every request, because the "
    "data lives here. For Thursday the browser plus a home-screen icon gives you the same "
    "phone test in minutes. After the 29th it can be a Tuesday project if the module "
    "wants a native build."
)
note(
    "Same Wi-Fi, laptop awake, server running, firewall rule present: those four are the "
    "causes, in order, when a phone page will not load."
)

# ---------------------------------------------------------------- 8
doc.add_heading("What was tested today", level=1)
bullets([
    "26 of 26 automated API checks passed, exit code 0, after two real bugs in the test itself were fixed (a wrong photo filename and an assertion that ran before its data existed).",
    "Resident login and staff login in two separate browser profiles, with the wrong password rejected.",
    "The full four-step report flow through the real interface: sample photo, location confirmed, Large, Deep, High, damage Yes, typed description, submit, reference #1025 on the details page with the reporter name Alex R. taken from the users table.",
    "Portal statistics of 98 total, 22 new, 33 fixed, 34 percent resolution, computed from the rows just written.",
    "The queue showing #1025 first because its score is the highest, with the priority badge visible on the details page.",
    "The status change to Verified with a real timestamp, and the resident's untouched browser tab showing the same change seconds later without a reload.",
    "Both database-generated notifications present: report received on submit, report verified on the staff action.",
    "Console clean on both tabs, npm run lint exit 0, npm run build exit 0, both rerun after the final edits.",
])

# ---------------------------------------------------------------- 9
doc.add_heading("Four problems solved today", level=1)
doc.add_paragraph("These are the challenges to mention if the panel asks what did not go smoothly.")
numbered([
    "The application login was rejected with login failed for user. The instance only allowed Windows authentication, so a SQL login could exist and still never connect. Enabling mixed-mode authentication and restarting the service resolved it.",
    "The SQL tools worked while the Node driver could not reach the server. sqlcmd talks over shared memory by default; the Node driver speaks TCP, and TCP was disabled on the instance. The fix was to enable TCP/IP on the loopback addresses with static port 1433, restarting the service again.",
    "The seeder failed inserting the synthetic rows with an identity error even though IDENTITY_INSERT was switched on inside a transaction. The switch and the inserts were not landing on the same database session. Sending both as a single batch resolved it, because one batch runs on one session.",
    "The API logged itself onto a random port. The process runner injects PORT=0 for a free port, and the server used that variable. It now reads its own variable, PP_PORT, defaults to 3001, and prints the port it actually bound.",
])
doc.add_paragraph(
    "Each one was found by running the thing and reading the error, not by guessing. That "
    "is also how the rest of today went: run it, look at the screen, fix the cause."
)

# ---------------------------------------------------------------- 10
doc.add_heading("The demo on Thursday 29 October", level=1)
doc.add_heading("Five minutes before", level=2)
numbered([
    "Start SQL Server if it is not already running (it is set to start automatically).",
    "npm run db:reset in server/, so the demo starts clean and the first submission is #1025.",
    "npm run build, then cd server && npm start. Open http://localhost:3001 in Chrome and in Edge.",
    "Sign in Chrome as alex@potholepal.test / resident123 and Edge as staff@potholepal.test / staff123.",
])
doc.add_heading("The walk-through", level=2)
numbered([
    "In Chrome, report a pothole: photo, confirm the location, Large, Deep, High, damage Yes, a short description, submit. Land on #1025 and point at the reference number and the timeline. Say: the reference comes from the database identity, not from the front end.",
    "Open the portal in Edge. Point at the four statistics and say they are counted from the same table, not typed in. Open the queue and show #1025 sitting first because of the priority score.",
    "Open #1025 in Edge and press Advance: Start review. The timeline stamps Verified.",
    "Back to Chrome, untouched until now. Within ten seconds the status updates on screen and the verified notification appears. Say: two browsers, one database, no page reload.",
    "Finish on the map screen with the full set of live markers, and mention the duplicate rule: try submitting the same spot twice and the system refuses with the existing reference instead of creating a second work item.",
])
doc.add_paragraph(
    "That sequence is criterion B2 in one pass: working functionality, architectural logic "
    "visible in the flow, no mockups on the critical path."
)

# ---------------------------------------------------------------- 11
doc.add_heading("Likely questions, short answers", level=1)
table(
    ["Question", "Answer"],
    [
        ["Why SQL Server?", "It is what we use in class, it was already on our machines, and the rubric asks for a justified stack rather than a named database."],
        ["How are passwords stored?", "As bcrypt hashes with 10 rounds. Plain text never touches the table. Everything else about security is deliberately post-demo."],
        ["How do duplicates work?", "Not fixed, within 50 metres, last 7 days. Bounding box first, then exact distance. The response refuses with the existing reference and the photo can be attached as evidence."],
        ["How does staff data reach the resident?", "Both sides read the same table through the same API. The front end refreshes the list every 10 seconds, so a status change appears on the other side without a reload."],
        ["Why is it not hosted?", "InfinityFree runs PHP and MySQL, not Node, and no free host covered Node plus a SQL database. Local hosting also removes internet and sleep risks during the demo."],
        ["What is mocked?", "The Teams crews, the monthly chart in Analytics, the badge artwork, and the phase-two points display. Reports, statuses, users, notifications, statistics and the queue are real."],
        ["How was it tested?", "26 automated API checks plus a manual two-browser walkthrough, all today. Lint and build both exit clean."],
        ["What happens if the server restarts?", "Sessions live in memory, so users sign in again. The data is in SQL Server and does not move."],
    ],
    widths=[1.9, 4.7],
)

# ---------------------------------------------------------------- 12
doc.add_heading("Still fake or unfinished", level=1)
bullets([
    "The Teams tab shows three sample crews; no crew table exists yet.",
    "The Analytics bar chart is a sample; the seeded data only reaches back 30 days, so there is nothing older to plot.",
    "Badge artwork and the Road Warrior label on the details screen are fixed text.",
    "Points are real in the database and update on fixed, but the Profile screen still shows its old layout around them.",
    "Sessions are in memory, so a server restart signs everyone out. Data is unaffected.",
    "The refresh is a poll every 10 seconds, not a live push.",
    "Reports nearby counts a bounding box around the town centre rather than a radius.",
    "No automated tests for the front end yet, and no error boundary: a screen that throws would take the page down.",
])

# ---------------------------------------------------------------- 13
doc.add_heading("Tomorrow, Wednesday 21 October", level=1)
numbered([
    "Commit tonight or first thing. There are 17 changed or new paths sitting uncommitted, including the whole server directory. That is the largest risk on this machine right now.",
    "Run the demo twice from scratch with db:reset between runs, once each of you driving, so nobody is surprised by a screen.",
    "Pull the smoke test output into the dissertation's evaluation section for criterion A4, and record the two-browser result as your system evaluation.",
    "Draw the architecture for the dissertation: browser, Express, SQL Server, uploads, with the submit flow as a sequence diagram. Criterion A3 wants UML diagrams and a justified stack.",
    "Wire the priority scores to the triage comparison figures due Thursday 22 October: the score and priority columns exist for exactly that comparison.",
    "Add the error boundary and decide whether the Monday rehearsal needs a recorded backup video of the demo.",
])

note(
    "The demo path in this document ran end to end today. npm run db:reset puts the data "
    "back to a clean start with the next report numbered 1025."
)

out = "docs/PotholePal-day2-backend.docx"
doc.save(out)
print("wrote", out)
